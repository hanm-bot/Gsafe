#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import argparse
import sys
import zipfile
from pathlib import Path
from lxml import etree

try:
    from docx import Document
    from docx.shared import Pt, Length
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
except ImportError:
    print("Thiếu python-docx. Cài: pip install python-docx", file=sys.stderr)
    sys.exit(3)

NS_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
NANG = "NẶNG"
NHE = "NHẸ"

def cap_heading(doan):
    ten = (doan.style.name or "") if doan.style is not None else ""
    if ten.startswith("Heading "):
        duoi = ten.split("Heading ", 1)[1].strip()
        if duoi.isdigit():
            return int(duoi)
    return None

def kiem_heading(tai_lieu):
    loi = []
    cap_truoc = 0
    for chi_so, doan in enumerate(tai_lieu.paragraphs, start=1):
        cap = cap_heading(doan)
        if cap is None:
            continue
        if cap_truoc and cap > cap_truoc + 1:
            loi.append((NANG, f"Heading nhảy cấp H{cap_truoc} → H{cap} tại đoạn {chi_so}: "
                              f"\"{doan.text.strip()[:60]}\""))
        cap_truoc = cap
    return loi

def kiem_bang(tai_lieu):
    loi = []
    for so_bang, bang in enumerate(tai_lieu.tables, start=1):
        if not bang.rows:
            loi.append((NANG, f"Bảng {so_bang}: không có hàng nào"))
            continue
        so_o_tieu_de = len(bang.rows[0].cells)
        for so_hang, hang in enumerate(bang.rows[1:], start=2):
            if len(hang.cells) != so_o_tieu_de:
                loi.append((NANG, f"Bảng {so_bang} hàng {so_hang}: {len(hang.cells)} ô, "
                                  f"hàng tiêu đề có {so_o_tieu_de} ô"))
    return loi

def kiem_font_tieng_viet(duong_dan):
    with zipfile.ZipFile(duong_dan) as z:
        ten = z.namelist()
        than = z.read("word/document.xml").decode("utf-8", errors="replace")
        kieu = z.read("word/styles.xml").decode("utf-8", errors="replace") if "word/styles.xml" in ten else ""
    if "w:eastAsia" in than or "w:eastAsia" in kieu:
        return []
    return [(NANG, "Không thấy khai báo w:eastAsia — Word có thể thay font và làm vỡ dấu tiếng Việt")]

def kiem_trang_trong(tai_lieu):
    loi = []
    truoc_la_ngat = False
    co_noi_dung_tu_lan_ngat = False
    for chi_so, doan in enumerate(tai_lieu.paragraphs, start=1):
        la_ngat = "w:br" in doan._p.xml and 'w:type="page"' in doan._p.xml
        if la_ngat:
            if truoc_la_ngat and not co_noi_dung_tu_lan_ngat:
                loi.append((NANG, f"Hai lần ngắt trang liền nhau, không có nội dung ở giữa (đoạn {chi_so}) "
                                  f"— sinh ra trang trống"))
            truoc_la_ngat = True
            co_noi_dung_tu_lan_ngat = False
            continue
        if doan.text.strip():
            co_noi_dung_tu_lan_ngat = True

    if truoc_la_ngat and not co_noi_dung_tu_lan_ngat:
        loi.append((NANG, "Ngắt trang ở cuối tài liệu, sau đó không có nội dung — sinh trang trống cuối"))
    return loi

def kiem_watermark(duong_dan, bat_buoc):
    with zipfile.ZipFile(duong_dan) as z:
        ten = z.namelist()
        header = [t for t in ten if t.startswith("word/header")]
        co_watermark = False
        for h in header:
            noi = z.read(h).decode("utf-8", errors="replace")
            if "WordArt" in noi or "PowerPlusWaterMarkObject" in noi or "watermark" in noi.lower():
                co_watermark = True
                break
        anh_header = [t for t in ten if t.startswith("word/media/")]
        
        theme_map = {}
        if "word/theme/theme1.xml" in ten:
            th = z.read("word/theme/theme1.xml")
            root = etree.fromstring(th)
            ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
            for font_scheme in root.xpath("//a:fontScheme", namespaces=ns):
                for font_type in ["majorFont", "minorFont"]:
                    f_elem = font_scheme.xpath(f"a:{font_type}", namespaces=ns)
                    if f_elem:
                        prefix = font_type.replace("Font", "")
                        for child in f_elem[0]:
                            tag = child.tag.split('}')[-1]
                            script = child.get("script", "")
                            font_val = child.get("typeface", "")
                            if tag == "latin":
                                theme_map[f"{prefix}HAnsi"] = font_val
                                theme_map[f"{prefix}Ascii"] = font_val
                            elif tag == "ea":
                                theme_map[f"{prefix}EastAsia"] = font_val
                            elif tag == "cs":
                                theme_map[f"{prefix}Bidi"] = font_val
                            elif tag == "font":
                                if script == "Hans":
                                    theme_map[f"{prefix}EastAsia"] = font_val
                                elif script == "Arab":
                                    theme_map[f"{prefix}Bidi"] = font_val
                                # Simplified: map everything to these 4 keys
                                

    if co_watermark:
        return [(NHE, f"Có watermark trong {len(header)} file header — đối chiếu bản gốc xem đúng trang chưa")], theme_map
    if bat_buoc:
        return [(NANG, "Bản gốc yêu cầu watermark nhưng file này KHÔNG có watermark trong header")], theme_map
    return [(NHE, f"Không thấy watermark (file có {len(anh_header)} ảnh nhúng) — "
                  "nếu bản gốc có watermark thì đã mất")], theme_map

def doc_numbering(duong_dan):
    num_to_abs = {}
    abs_to_ilvl = {}
    with zipfile.ZipFile(duong_dan) as z:
        if "word/numbering.xml" in z.namelist():
            num_xml = z.read("word/numbering.xml")
            root = etree.fromstring(num_xml)
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            for num in root.xpath("//w:num", namespaces=ns):
                num_id = num.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numId")
                abs_elem = num.xpath("w:abstractNumId", namespaces=ns)
                if abs_elem:
                    num_to_abs[num_id] = abs_elem[0].get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val")
                    
            for abs_num in root.xpath("//w:abstractNum", namespaces=ns):
                abs_id = abs_num.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}abstractNumId")
                abs_to_ilvl[abs_id] = {}
                for lvl in abs_num.xpath("w:lvl", namespaces=ns):
                    ilvl = lvl.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ilvl")
                    fonts = lvl.xpath("w:rPr/w:rFonts", namespaces=ns)
                    font_dict = {}
                    if fonts:
                        f = fonts[0]
                        for k in ["ascii", "hAnsi", "eastAsia", "cs", "asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"]:
                            val = f.get(f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{k}")
                            if val: font_dict[k] = val
                    abs_to_ilvl[abs_id][ilvl] = font_dict
    return num_to_abs, abs_to_ilvl

def kiem_nd30_tang_1(tai_lieu, theme_map, num_to_abs, abs_to_ilvl):
    loi = []
    
    # Extract styles font to resolve inheritance
    styles_font = {}
    for st in tai_lieu.styles:
        if st.type == 1: # paragraph style
            rPr = st._element.rPr
            st_fonts = {}
            if rPr is not None and rPr.rFonts is not None:
                for k in ["ascii", "hAnsi", "eastAsia", "cs", "asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"]:
                    val = rPr.rFonts.get(qn(f"w:{k}"))
                    if val: st_fonts[k] = val
            styles_font[st.name] = st_fonts
    
    # Helpers
    def get_eff_fonts(run, para):
        eff = {"ascii": "Times New Roman", "hAnsi": "Times New Roman", "eastAsia": "Times New Roman", "cs": "Times New Roman"}
        
        # 1. Base from Normal style
        norm = styles_font.get("Normal", {})
        
        # 2. Paragraph style
        st_name = para.style.name if para.style else "Normal"
        p_st = styles_font.get(st_name, norm)
        
        # 3. Run properties
        r_fonts = {}
        rPr = run._element.rPr
        if rPr is not None and rPr.rFonts is not None:
            for k in ["ascii", "hAnsi", "eastAsia", "cs", "asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"]:
                val = rPr.rFonts.get(qn(f"w:{k}"))
                if val: r_fonts[k] = val
        
        for k in ["ascii", "hAnsi", "eastAsia", "cs"]:
            # Check Run explicitly
            if f"{k}Theme" in r_fonts:
                eff[k] = theme_map.get(r_fonts[f"{k}Theme"], r_fonts[f"{k}Theme"])
            elif k in r_fonts:
                eff[k] = r_fonts[k]
            # Else check Style
            elif f"{k}Theme" in p_st:
                eff[k] = theme_map.get(p_st[f"{k}Theme"], p_st[f"{k}Theme"])
            elif k in p_st:
                eff[k] = p_st[k]
            # Else check Normal style
            elif f"{k}Theme" in norm:
                eff[k] = theme_map.get(norm[f"{k}Theme"], norm[f"{k}Theme"])
            elif k in norm:
                eff[k] = norm[k]
                
            # If still theme (e.g. implicitly set from docDefaults)
            if eff[k] in theme_map:
                eff[k] = theme_map[eff[k]]
                
        return eff
        
    def get_eff_size(run, para):
        if run.font and run.font.size: return run.font.size.pt
        if para.style and para.style.font and para.style.font.size: return para.style.font.size.pt
        st_name = para.style.name if para.style else "Normal"
        # Not implementing full size inheritance, just assume 13 if missing
        return 13

    # Section checks (6.1, 6.2, 6.8)
    total_pages = 2 # approximation, 6.8 ignores if 1 page, we just assume >1 for test
    for i, sec in enumerate(tai_lieu.sections):
        w = sec.page_width.mm if sec.page_width else 0
        h = sec.page_height.mm if sec.page_height else 0
        
        # 6.1
        if w > h:
            loi.append((NHE, f"Section {i+1}: hướng ngang — đối chiếu ngoại lệ bảng biểu"))
        elif not (209 <= w <= 211 and 296 <= h <= 298):
            loi.append((NANG, f"Section {i+1}: khổ giấy {w:.0f}x{h:.0f}mm không phải A4"))
            
        # 6.2
        t = sec.top_margin.mm if sec.top_margin else 0
        b = sec.bottom_margin.mm if sec.bottom_margin else 0
        l = sec.left_margin.mm if sec.left_margin else 0
        r = sec.right_margin.mm if sec.right_margin else 0
        
        if not (19.5 <= t <= 25.5 and 19.5 <= b <= 25.5 and 29.5 <= l <= 35.5 and 14.5 <= r <= 20.5):
            loi.append((NANG, f"Section {i+1}: lề sai {t:.1f}/{b:.1f}/{l:.1f}/{r:.1f}mm"))
            
        # 6.8
        # Since finding PAGE in header is complex with docx, we just check the XML
        has_page = False
        header_center = False
        first_page_ok = sec.different_first_page_header_footer
        
        for hdr in [sec.header, sec.first_page_header, sec.even_page_header]:
            if hdr and not hdr.is_linked_to_previous:
                for p in hdr.paragraphs:
                    if 'PAGE' in p._p.xml:
                        has_page = True
                        if p.alignment == WD_ALIGN_PARAGRAPH.CENTER or 'jc w:val="center"' in p._p.xml:
                            header_center = True
        
        if not first_page_ok:
            loi.append((NANG, f"Section {i+1}: trang đầu có số (chưa bật different_first_page)"))
        if not has_page:
            loi.append((NANG, f"Section {i+1}: không có số trang (trường PAGE)"))

    # Paragraph checks
    for p_idx, p in enumerate(tai_lieu.paragraphs, start=1):
        txt = p.text.strip()
        if not txt: continue
        preview = txt[:30]
        st_name = p.style.name if p.style else ""
        
        # 6.6, 6.7, 6.9
        is_normal = (st_name == "Normal" or not st_name) and not p._p.xpath('ancestor::w:tbl')
        
        if is_normal:
            # 6.6 Giãn dòng
            ls = p.paragraph_format.line_spacing
            if ls is None and p.style and hasattr(p.style, "paragraph_format"):
                ls = p.style.paragraph_format.line_spacing
            if ls is not None and ls > 1.5:
                loi.append((NANG, f"Đoạn {p_idx} ({preview}): giãn dòng > 1.5"))
            
            # 6.7 Khoảng cách đoạn
            sb = p.paragraph_format.space_before.pt if p.paragraph_format.space_before is not None else None
            sa = p.paragraph_format.space_after.pt if p.paragraph_format.space_after is not None else None
            if sb is None and p.style and hasattr(p.style, "paragraph_format"):
                sb = p.style.paragraph_format.space_before.pt if p.style.paragraph_format.space_before is not None else 0
            if sa is None and p.style and hasattr(p.style, "paragraph_format"):
                sa = p.style.paragraph_format.space_after.pt if p.style.paragraph_format.space_after is not None else 0
            sb = sb if sb is not None else 0
            sa = sa if sa is not None else 0
            if sb + sa < 6:
                loi.append((NANG, f"Đoạn {p_idx} ({preview}): khoảng cách đoạn < 6pt"))
                
            # 6.9 Canh lề
            al = p.alignment
            if al is None and p.style and hasattr(p.style, "paragraph_format"):
                al = p.style.paragraph_format.alignment
            if al != WD_ALIGN_PARAGRAPH.JUSTIFY:
                loi.append((NHE, f"Đoạn {p_idx} ({preview}): chưa canh đều hai bên (justify)"))
                
        # 6.3 mở rộng: kiểm tra numbering font
        num_pr = p._p.pPr.numPr if p._p.pPr is not None else None
        if num_pr is None and p.style and p.style.type == 1:
            if p.style._element.pPr is not None:
                num_pr = p.style._element.pPr.numPr
        
        if num_pr is not None:
            num_id = num_pr.xpath("w:numId/@w:val")
            ilvl = num_pr.xpath("w:ilvl/@w:val")
            if num_id:
                nid = num_id[0]
                ilv = ilvl[0] if ilvl else "0"
                abs_id = num_to_abs.get(nid)
                if abs_id:
                    num_fonts = abs_to_ilvl.get(abs_id, {}).get(ilv, {})
                    # Giải mã theme
                    eff_num = {}
                    for k in ["ascii", "hAnsi", "eastAsia", "cs"]:
                        if f"{k}Theme" in num_fonts:
                            eff_num[k] = theme_map.get(num_fonts[f"{k}Theme"], num_fonts[f"{k}Theme"])
                        elif k in num_fonts:
                            eff_num[k] = num_fonts[k]
                    
                    for script, f_name in eff_num.items():
                        if f_name != "Times New Roman":
                            loi.append((NANG, f"Đoạn {p_idx} ({preview}): ký tự đầu dòng (numId={nid}, ilvl={ilv}) dùng phông {f_name} khác Times New Roman"))
                            break
        
        for r_idx, r in enumerate(p.runs):
            if not r.text.strip(): continue
            
            # 6.3 Font
            fonts = get_eff_fonts(r, p)
            if st_name == "NĐ30 Khối mã":
                pass
            else:
                for script, f_name in fonts.items():
                    if f_name != "Times New Roman":
                        loi.append((NANG, f"Đoạn {p_idx} ({preview}): sai phông chữ {f_name} (ở {script})"))
                        break
                        
            # 6.10 Emoji
            for idx, c in enumerate(r.text):
                code = ord(c)
                if (0x2300 <= code <= 0x23FF) or (0x2600 <= code <= 0x27BF) or (0x2B00 <= code <= 0x2BFF) or (0x1F000 <= code <= 0x1FAFF) or (code == 0xFE0F):
                    loi.append((NANG, f"Đoạn {p_idx} ({preview}): có emoji/ký hiệu '{c}' tại vị trí {idx}"))
            
            # 6.4 Color
            # python-docx rgb is None if auto, or RGBColor. If it's a theme color, it's not RGBColor.
            # We just check RGB for non-black/non-auto
            color = r.font.color.rgb if r.font and r.font.color else None
            if color is not None and str(color) not in ["000000", ""]:
                loi.append((NANG, f"Đoạn {p_idx} ({preview}): chữ màu {color} (phải đen/tự động)"))
            
            # check shading for white text on colored bg
            shd = r._element.xpath('.//w:shd/@w:fill')
            if shd and shd[0] not in ['auto', '000000', 'FFFFFF'] and str(color) == 'FFFFFF':
                loi.append((NANG, f"Đoạn {p_idx} ({preview}): chữ trắng trên nền màu"))
                
            # 6.5 Size
            size = get_eff_size(r, p)
            if size < 11:
                loi.append((NANG, f"Đoạn {p_idx} ({preview}): cỡ chữ {size} < 11"))
            elif is_normal and size < 13:
                loi.append((NANG, f"Đoạn {p_idx} ({preview}): nội dung cỡ {size} < 13"))
                
    # Check tables
    for t_idx, t in enumerate(tai_lieu.tables, start=1):
        for r_idx, row in enumerate(t.rows):
            for c_idx, cell in enumerate(row.cells):
                for p in cell.paragraphs:
                    for run in p.runs:
                        if not run.text.strip(): continue
                        color = run.font.color.rgb if run.font and run.font.color else None
                        if color is not None and str(color) == "FFFFFF":
                            loi.append((NANG, f"Bảng {t_idx}: có chữ trắng (ô hàng {r_idx+1})"))

    return loi


def main():
    for luong in (sys.stdout, sys.stderr):
        if hasattr(luong, "reconfigure"):
            luong.reconfigure(encoding="utf-8", errors="replace")

    p = argparse.ArgumentParser(description="Kiểm .docx trước khi phát hành")
    p.add_argument("docx", help="File .docx cần kiểm")
    p.add_argument("--watermark-bat-buoc", action="store_true",
                   help="Bản gốc có watermark — thiếu là lỗi NẶNG")
    tham_so = p.parse_args()

    duong_dan = Path(tham_so.docx)
    if not duong_dan.exists():
        print(f"Không thấy file: {duong_dan}", file=sys.stderr)
        return 3

    tai_lieu = Document(str(duong_dan))
    loi = []
    loi += kiem_heading(tai_lieu)
    loi += kiem_bang(tai_lieu)
    loi += kiem_font_tieng_viet(duong_dan)
    loi += kiem_trang_trong(tai_lieu)
    loi_watermark, theme_map = kiem_watermark(duong_dan, tham_so.watermark_bat_buoc)
    loi += loi_watermark
    num_to_abs, abs_to_ilvl = doc_numbering(duong_dan)
    loi += kiem_nd30_tang_1(tai_lieu, theme_map, num_to_abs, abs_to_ilvl)

    nang = [m for m in loi if m[0] == NANG]
    nhe = [m for m in loi if m[0] == NHE]

    print(f"KIỂM XUẤT BẢN: {duong_dan.name}")
    print(f"  Đoạn: {len(tai_lieu.paragraphs)} · Bảng: {len(tai_lieu.tables)}")
    print()

    for muc, mo_ta in nang:
        print(f"  [{muc}] {mo_ta}")
    for muc, mo_ta in nhe:
        print(f"  [{muc}] {mo_ta}")

    print()
    if nang:
        print(f"❌ CHẶN PHÁT HÀNH — {len(nang)} lỗi nặng phải sửa trước.")
        return 2
    print(f"✅ Đạt phần máy kiểm được ({len(nhe)} ghi chú cần người soát đối chiếu).")
    print("   Lưu ý: máy không kiểm được thể thức nghiệp vụ và giọng văn.")
    return 0

if __name__ == "__main__":
    sys.exit(main())
